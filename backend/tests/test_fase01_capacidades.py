"""SEC-01 — a política financeira vale igual em todos os canais.

Matriz perfil × canal. Um título sintético com marcador único é criado; o
operacional não pode recebê-lo por nenhum formato (JSON, CSV, XLSX, PDF),
nem em corpo nem em metadado. Admin e financeiro continuam recebendo.

Antes da correção: operacional levava 403 em /billings e /reports mas 200
com o marcador em /exports/billings e no PDF da linha do tempo.
"""
from __future__ import annotations

import io
from datetime import date
from decimal import Decimal

import pytest
from fastapi.testclient import TestClient
from pypdf import PdfReader

from app.api.deps import get_current_user
from app.core.permissions import Capability, has_capability, roles_with
from app.db.session import get_db
from app.main import app
from app.models.billing import Billing
from app.models.enums import BillingStatus, UserRole
from app.models.user import User

MARCADOR = "MARCADOR-FIN-7731"
VALOR = Decimal("7731.42")
VALOR_TXT = "7731.42"
VALOR_BR = "7.731,42"

ROLES = {
    "admin": UserRole.ADMIN,
    "operacional": UserRole.OPERATIONAL,
    "financeiro": UserRole.FINANCIAL,
    "cliente": UserRole.CLIENT,
}


@pytest.fixture()
def titulo(db, contrato):
    b = Billing(
        contract_id=contrato.id,
        client_id=contrato.client_id,
        amount=VALOR,
        paid_amount=None,
        due_date=date(2020, 3, 10),
        status=BillingStatus.OVERDUE,
        billing_type="avulsa",
        title=MARCADOR,
    )
    db.add(b)
    db.commit()
    db.refresh(b)
    return b


@pytest.fixture()
def como(db):
    """como('operacional') → TestClient autenticado com esse perfil.
    None → anônimo. Um override por vez (as fixtures http_* compartilham)."""

    def _factory(perfil: str | None) -> TestClient:
        app.dependency_overrides[get_db] = lambda: db
        if perfil is None:
            app.dependency_overrides.pop(get_current_user, None)
        else:
            user = User(
                id=90, name=f"Perfil {perfil}", email=f"{perfil}@cap.test",
                role=ROLES[perfil], active=True, is_deleted=False, password_hash="x",
            )
            app.dependency_overrides[get_current_user] = lambda: user
        return TestClient(app, raise_server_exceptions=False)

    yield _factory
    app.dependency_overrides.clear()


def _texto(resp) -> str:
    """Conteúdo legível da resposta, qualquer formato."""
    ctype = resp.headers.get("content-type", "")
    meta = " ".join(f"{k}:{v}" for k, v in resp.headers.items())
    if "pdf" in ctype:
        reader = PdfReader(io.BytesIO(resp.content))
        corpo = "\n".join(p.extract_text() or "" for p in reader.pages)
        corpo += " " + " ".join(str(v) for v in (reader.metadata or {}).values())
        return corpo + " " + meta
    if "spreadsheet" in ctype:
        from openpyxl import load_workbook
        wb = load_workbook(io.BytesIO(resp.content))
        celulas = [str(c) for ws in wb.worksheets for row in ws.iter_rows(values_only=True) for c in row if c is not None]
        return " ".join(celulas) + " " + meta
    return resp.text + " " + meta


def _tem_financeiro(texto: str) -> bool:
    return MARCADOR in texto or VALOR_TXT in texto or VALOR_BR in texto


# Canais que entregam o título inteiro: só quem tem FINANCIAL_READ.
CANAIS_FINANCEIROS = [
    "/api/v1/exports/billings?fmt=csv",
    "/api/v1/exports/billings?fmt=xlsx",
    "/api/v1/exports/billings-report?fmt=csv&situacao=todas",
    "/api/v1/exports/billings-report?fmt=xlsx&situacao=todas",
    "/api/v1/exports/billings-report?fmt=pdf&situacao=todas",
    "/api/v1/exports/delinquents?fmt=csv",
    "/api/v1/exports/delinquents?fmt=xlsx",
    "/api/v1/exports/delinquents?fmt=pdf",
    "/api/v1/billings/",
    "/api/v1/billings/exports/csv",
    "/api/v1/reports/summary",
]


class TestMatrizCapacidades:
    def test_matriz_declarada(self):
        assert roles_with(Capability.FINANCIAL_READ) == (UserRole.ADMIN, UserRole.FINANCIAL)
        assert has_capability("operacional", Capability.REGISTRY_READ)
        assert not has_capability("operacional", Capability.FINANCIAL_READ)
        assert not has_capability("cliente", Capability.REGISTRY_READ)
        assert not has_capability("inexistente", Capability.FINANCIAL_READ)
        assert not has_capability(None, Capability.FINANCIAL_READ)


class TestCanaisFinanceiros:
    @pytest.mark.parametrize("url", CANAIS_FINANCEIROS)
    @pytest.mark.parametrize("perfil,esperado", [
        ("operacional", 403), ("cliente", 403), (None, 401),
    ])
    def test_sem_capacidade_nao_recebe(self, como, titulo, url, perfil, esperado):
        r = como(perfil).get(url)
        assert r.status_code == esperado, r.text[:200]
        assert not _tem_financeiro(_texto(r))

    @pytest.mark.parametrize("url", [u for u in CANAIS_FINANCEIROS if "exports" in u and "reports/" not in u])
    @pytest.mark.parametrize("perfil", ["admin", "financeiro"])
    def test_com_capacidade_recebe_o_titulo(self, como, titulo, url, perfil):
        r = como(perfil).get(url)
        assert r.status_code == 200, r.text[:200]
        assert _tem_financeiro(_texto(r)), f"{url} não trouxe o título para {perfil}"


class TestTimelinePdf:
    def _url(self, cliente):
        return f"/api/v1/clients/{cliente.id}/timeline-pdf"

    def test_operacional_recebe_pdf_sem_contratos_nem_cobrancas(self, como, cliente, titulo, ordem_servico):
        r = como("operacional").get(self._url(cliente))
        assert r.status_code == 200
        texto = _texto(r)
        assert not _tem_financeiro(texto)
        assert "COBRANÇAS" not in texto and "CONTRATOS" not in texto
        # o conteúdo operacional continua lá
        assert "OS-2025-001" in texto
        assert "omitidos" in texto

    @pytest.mark.parametrize("perfil", ["admin", "financeiro"])
    def test_financeiro_e_admin_recebem_cobrancas(self, como, cliente, titulo, perfil):
        r = como(perfil).get(self._url(cliente))
        assert r.status_code == 200
        texto = _texto(r)
        assert _tem_financeiro(texto)
        assert "CONTRATOS" in texto

    @pytest.mark.parametrize("perfil,esperado", [("cliente", 403), (None, 401)])
    def test_negado(self, como, cliente, titulo, perfil, esperado):
        assert como(perfil).get(self._url(cliente)).status_code == esperado

    def test_json_e_pdf_aplicam_a_mesma_regra(self, como, cliente, titulo):
        tc = como("operacional")
        js = tc.get(f"/api/v1/clients/{cliente.id}/timeline", params={"category": "financeiro"})
        assert js.status_code == 200 and js.json()["items"] == []
        assert not _tem_financeiro(_texto(tc.get(self._url(cliente))))


class TestDashboard:
    def test_operacional_sem_bloco_financeiro(self, como, titulo):
        r = como("operacional").get("/api/v1/dashboard/")
        assert r.status_code == 200
        body = r.json()
        assert body["finance"] is None
        assert body["upcoming_billings"] == []
        assert not _tem_financeiro(r.text)
        # contagens de cadastro/campo continuam
        assert body["clients"]["active"] >= 1
        assert "service_orders" in body and "trackers" in body

    @pytest.mark.parametrize("perfil", ["admin", "financeiro"])
    def test_financeiro_mantem_contrato(self, como, titulo, perfil):
        body = como(perfil).get("/api/v1/dashboard/").json()
        assert set(body["finance"]) == {
            "pending_count", "overdue_count", "received_month",
            "received_prev_month", "delta_received", "delta_pct",
        }
        assert body["finance"]["overdue_count"] >= 1


class TestDelinquencyStatus:
    def test_operacional_ve_so_contagem_de_clientes(self, como, titulo):
        body = como("operacional").get("/api/v1/delinquency/status").json()
        assert body["valor_total_vencido"] is None
        assert body["cobrancas_vencidas"] is None
        assert isinstance(body["clientes_inadimplentes"], int)

    @pytest.mark.parametrize("perfil", ["admin", "financeiro"])
    def test_financeiro_ve_valor(self, como, titulo, perfil):
        body = como(perfil).get("/api/v1/delinquency/status").json()
        assert body["valor_total_vencido"] == pytest.approx(float(VALOR))
        assert body["cobrancas_vencidas"] == 1


class TestExportsCadastraisPreservados:
    @pytest.mark.parametrize("url", [
        "/api/v1/exports/clients?fmt=csv",
        "/api/v1/exports/vehicles?fmt=xlsx",
        "/api/v1/exports/trackers?fmt=csv",
    ])
    def test_operacional_continua_exportando_cadastro(self, como, titulo, url):
        r = como("operacional").get(url)
        assert r.status_code == 200
        assert not _tem_financeiro(_texto(r))
