"""
Testes de integração para /api/v1/billing-closure.

Cobertos:
- GET /simulate          → estrutura da resposta, uninstall_events presentes,
                           filtros pf/pj/client, mês ausente → 422,
                           client_id ausente quando filter_type=client → 422
- GET /simulate/pdf      → content-type application/pdf, bytes não vazios, header %PDF
- POST /generate         → execução síncrona, retorna status='completed' imediatamente
                           com campos generated/total_amount/grand_total,
                           mês ausente → 422, client_id ausente → 422
- Autorização:
    OPERATIONAL → 403 em todos os endpoints
    FINANCIAL   → 200 em todos os endpoints
    sem auth    → 401
"""
from __future__ import annotations

from datetime import date
from decimal import Decimal
from io import BytesIO

import pytest
from openpyxl import load_workbook
from sqlalchemy import select

from app.models.billing import Billing
from app.models.closure_job import ClosureJob
from app.models.client_charge_item import ClientChargeItem
from app.models.contract import Contract
from app.models.enums import BillingStatus
from app.models.uninstall_event import UninstallEvent

PREFIX = "/api/v1/billing-closure"
REF_MONTH = "2025-05"


def test_september_service_closure_is_due_in_october(
    http, db, cliente, veiculo, contrato,
):
    event = UninstallEvent(
        vehicle_id=veiculo.id, contract_id=contrato.id, client_id=cliente.id,
        uninstall_date=date(2026, 9, 10), fee_amount=Decimal('160.00'), status='pending',
    )
    future_event = UninstallEvent(
        vehicle_id=veiculo.id, client_id=cliente.id,
        uninstall_date=date(2026, 10, 10), fee_amount=Decimal('50.00'), status='pending',
    )
    charge = ClientChargeItem(
        client_id=cliente.id, contract_id=contrato.id, title='Serviço de setembro',
        quantity=1, unit_price=Decimal('30.00'), total_amount=Decimal('30.00'),
        installment_count=1, start_date=date(2026, 9, 20), active=True,
    )
    db.add_all([event, future_event, charge])
    db.commit()

    params = {'service_month': '2026-09'}
    preview = http.get(PREFIX + '/simulate', params=params)
    assert preview.status_code == 200, preview.text
    data = preview.json()
    assert data['reference_month'] == '09/2026'
    assert data['billing_month'] == '10/2026'
    monthly = next(item for item in data['items'] if item['contract_id'] == contrato.id)
    assert monthly['due_date'] == '2026-10-15'
    assert monthly['service_period_label'] == '09/2026'
    assert [item['event_id'] for item in data['uninstall_events']] == [event.id]
    assert [item['item_id'] for item in data['charge_items']] == [charge.id]

    export = http.get(PREFIX + '/simulate/xlsx', params=params)
    assert export.status_code == 200, export.text
    sheet = load_workbook(BytesIO(export.content), read_only=True)['Contratos']
    assert sheet.cell(2, 12).value.date() == date(2026, 10, 15)
    assert sheet.cell(2, 13).value == '09/2026'

    result = http.post(PREFIX + '/generate', params={
        **params, 'contract_ids': contrato.id,
        'uninstall_event_ids': event.id, 'charge_item_ids': charge.id,
    })
    assert result.status_code == 200, result.text
    assert result.json()['reference_month'] == '2026-09'
    assert result.json()['billing_month'] == '10/2026'
    lote = db.get(ClosureJob, result.json()['closure_batch_id'])
    assert lote.reference_month == '2026-09'
    assert lote.result['payment_billing_ids'] == result.json()['payment_billing_ids']
    billings = db.scalars(select(Billing).where(
        Billing.id.in_(
            result.json()['billing_ids']
            + result.json()['uninstall_billing_ids']
            + result.json()['service_billing_ids']
        )
    )).all()
    assert len(billings) == 3
    assert {billing.due_date for billing in billings} == {date(2026, 10, 15)}

    visible = http.get('/api/v1/billings/', params={'client_id': cliente.id})
    assert visible.status_code == 200, visible.text
    visible_ids = {entry['id'] for entry in visible.json()}
    assert set(result.json()['payment_billing_ids']) <= visible_ids
    assert not {billing.id for billing in billings if billing.substituted_by_id} & visible_ids
    audit = http.get('/api/v1/billings/', params={
        'client_id': cliente.id, 'include_substituted': True,
    })
    assert audit.status_code == 200, audit.text
    assert {billing.id for billing in billings} <= {entry['id'] for entry in audit.json()}

    again = http.get(PREFIX + '/simulate', params=params)
    assert again.json()['to_generate'] == 0
    assert again.json()['uninstall_events'] == []
    assert again.json()['charge_items'] == []


def test_recibo_e_nfse_do_boleto_unico_discriminam_servicos_de_setembro(
    http, db, cliente, plan, veiculo, monkeypatch,
):
    from pypdf import PdfReader

    from app.services.billing_closure import execute_closure
    from app.services import nfse_provider

    contracts = [Contract(
        client_id=cliente.id, plan_id=plan.id, vehicle_id=veiculo.id,
        start_date=date(2026, 1, 1), status='ativo', billing_day=15,
    ) for _ in range(4)]
    db.add_all(contracts)
    db.commit()

    result = execute_closure(db, date(2026, 10, 1), activity_month=date(2026, 9, 1))
    assert result['payment_titles_generated'] == 1
    parent = db.get(Billing, result['payment_billing_ids'][0])
    assert parent.billing_type == 'boleto_unico'
    assert parent.period_label == '09/2026'
    assert parent.title.startswith('Fechamento 09/2026')
    assert parent.due_date == date(2026, 10, 15)
    components = http.get(f'/api/v1/billings/{parent.id}/components')
    assert components.status_code == 200, components.text
    assert len(components.json()) == 4
    assert sum((Decimal(str(item['amount'])) for item in components.json()), Decimal('0.00')) == parent.amount
    assert {item['substituted_by_id'] for item in components.json()} == {parent.id}

    description = nfse_provider.descricao_fechamento(db, parent)
    assert 'competência 09/2026' in description
    assert description.count('PLACA:') == 4
    assert description.count('R$ 99,90') == 4
    captured = {}
    class Provider:
        @staticmethod
        def emitir_nfse(*args, **kwargs):
            captured.update(kwargs)
            return 'simulada'
    monkeypatch.setattr(nfse_provider, 'modulo', lambda: Provider)
    assert nfse_provider.emitir_nfse(db, parent, cliente) == 'simulada'
    assert captured['discriminacao'] == description
    assert captured['competencia'] is None  # data oficial fica a da emissão

    parent.status = BillingStatus.PAID
    parent.payment_date = date(2026, 10, 15)
    parent.paid_amount = parent.amount
    db.commit()
    response = http.get(f'/api/v1/billings/{parent.id}/receipt')
    assert response.status_code == 200, response.text
    pages = PdfReader(BytesIO(response.content)).pages
    assert len(pages) == 2
    first = pages[0].extract_text()
    detail = pages[1].extract_text()
    assert 'MAIS 2 ITENS' in first
    assert 'PLACA' in first
    assert 'REF. 09/2026' in first
    assert 'FECHAMENTO 09/2026 - BOLETO' not in first
    assert 'DETALHAMENTO DO RECIBO' in detail
    assert detail.count('REF. 09/2026') == 4


def test_service_month_includes_first_and_final_prorata_in_next_month(
    http, db, cliente, plan,
):
    installed = Contract(
        client_id=cliente.id, plan_id=plan.id, start_date=date(2026, 9, 11),
        status='ativo', billing_day=10,
    )
    uninstalled = Contract(
        client_id=cliente.id, plan_id=plan.id, start_date=date(2026, 1, 1),
        end_date=date(2026, 9, 10), uninstalled_at=date(2026, 9, 10),
        status='cancelado', billing_day=10,
    )
    db.add_all([installed, uninstalled])
    db.commit()

    response = http.get(PREFIX + '/simulate', params={'service_month': '2026-09'})
    assert response.status_code == 200, response.text
    rows = {item['contract_id']: item for item in response.json()['items']}
    assert rows[installed.id]['due_date'] == '2026-10-10'
    assert rows[installed.id]['prorated_days'] == 20
    assert rows[uninstalled.id]['due_date'] == '2026-10-10'
    assert rows[uninstalled.id]['prorated_days'] == 10


def test_december_service_month_is_due_in_january(http, db, cliente, plan):
    contract = Contract(
        client_id=cliente.id, plan_id=plan.id, start_date=date(2026, 1, 1),
        status='ativo', billing_day=15,
    )
    db.add(contract)
    db.commit()

    response = http.get(PREFIX + '/simulate', params={'service_month': '2026-12'})
    assert response.status_code == 200, response.text
    assert response.json()['billing_month'] == '01/2027'
    assert response.json()['items'][0]['due_date'] == '2027-01-15'


def test_service_month_keeps_contract_ending_in_service_month(http, db, cliente, plan):
    contract = Contract(
        client_id=cliente.id, plan_id=plan.id, start_date=date(2026, 1, 1),
        end_date=date(2026, 9, 30), status='ativo', billing_day=15,
    )
    db.add(contract)
    db.commit()

    response = http.get(PREFIX + '/simulate', params={'service_month': '2026-09'})
    assert response.status_code == 200, response.text
    assert response.json()['items'][0]['due_date'] == '2026-10-15'


def test_remaining_service_installment_starts_in_next_month(http, db, cliente, contrato):
    charge = ClientChargeItem(
        client_id=cliente.id, contract_id=contrato.id, title='Serviço parcelado',
        quantity=1, unit_price=Decimal('60.00'), total_amount=Decimal('60.00'),
        installment_count=2, start_date=date(2026, 9, 20), active=True,
    )
    db.add(charge)
    db.flush()
    db.add(Billing(
        client_id=cliente.id, contract_id=contrato.id, item_id=charge.id,
        title='Serviço parcelado 1/2', billing_type='item', amount=Decimal('30.00'),
        due_date=date(2026, 9, 20), period_label='09/2026',
        installment_number=1, installment_total=2, status=BillingStatus.PAID,
    ))
    db.commit()

    result = http.post(PREFIX + '/generate', params={
        'service_month': '2026-09', 'contract_ids': -1,
        'uninstall_event_ids': -1, 'charge_item_ids': charge.id,
    })
    assert result.status_code == 200, result.text
    assert result.json()['services_generated'] == 1
    generated = db.get(Billing, result.json()['service_billing_ids'][0])
    assert generated.installment_number == 2
    assert generated.due_date == date(2026, 10, 15)


# ---------------------------------------------------------------------------
# GET /simulate
# ---------------------------------------------------------------------------

class TestSimulate:
    def test_requires_reference_month(self, http):
        r = http.get(PREFIX + "/simulate")
        assert r.status_code == 422

    def test_returns_200_with_month(self, http):
        r = http.get(PREFIX + "/simulate", params={"reference_month": REF_MONTH})
        assert r.status_code == 200

    def test_response_structure(self, http):
        r = http.get(PREFIX + "/simulate", params={"reference_month": REF_MONTH})
        data = r.json()
        for key in (
            "reference_month", "total_contracts", "to_generate", "already_generated",
            "items", "uninstall_events", "total_uninstall_fees", "grand_total",
        ):
            assert key in data

    def test_reference_month_value_in_response(self, http):
        r = http.get(PREFIX + "/simulate", params={"reference_month": REF_MONTH})
        assert r.json()["reference_month"] == "05/2025"

    def test_invalid_month_format_returns_422(self, http):
        r = http.get(PREFIX + "/simulate", params={"reference_month": "05-2025"})
        assert r.status_code == 422

    def test_invalid_month_value_returns_422(self, http):
        r = http.get(PREFIX + "/simulate", params={"reference_month": "2025-13"})
        assert r.status_code == 422

    def test_filter_client_without_client_id_returns_422(self, http):
        r = http.get(PREFIX + "/simulate", params={
            "reference_month": REF_MONTH,
            "filter_type": "client",
        })
        assert r.status_code == 422

    def test_filter_pf_returns_200(self, http):
        r = http.get(PREFIX + "/simulate", params={
            "reference_month": REF_MONTH,
            "filter_type": "pf",
        })
        assert r.status_code == 200

    def test_filter_pj_returns_200(self, http):
        r = http.get(PREFIX + "/simulate", params={
            "reference_month": REF_MONTH,
            "filter_type": "pj",
        })
        assert r.status_code == 200

    def test_uninstall_event_appears_in_response(self, http, db, cliente, veiculo):
        e = UninstallEvent(
            vehicle_id=veiculo.id,
            client_id=cliente.id,
            uninstall_date=date(2025, 5, 10),
            fee_amount=Decimal("100.00"),
            status="pending",
        )
        db.add(e)
        db.commit()
        r = http.get(PREFIX + "/simulate", params={"reference_month": REF_MONTH})
        assert r.status_code == 200
        data = r.json()
        assert len(data["uninstall_events"]) >= 1
        item = data["uninstall_events"][0]
        assert "event_id" in item
        assert "fee_amount" in item
        assert "skipped" in item
        assert item["client_id"] == cliente.id

    def test_skipped_event_appears_with_reason(self, http, db, cliente, veiculo):
        e = UninstallEvent(
            vehicle_id=veiculo.id,
            client_id=cliente.id,
            uninstall_date=date(2025, 5, 10),
            fee_amount=Decimal("1.00"),
            status="pending",
        )
        db.add(e)
        db.commit()
        r = http.get(PREFIX + "/simulate", params={"reference_month": REF_MONTH})
        assert r.status_code == 200
        skipped = [e for e in r.json()["uninstall_events"] if e["skipped"]]
        assert len(skipped) >= 1
        assert skipped[0]["skip_reason"] is not None

    def test_operational_cannot_access(self, http_op):
        r = http_op.get(PREFIX + "/simulate", params={"reference_month": REF_MONTH})
        assert r.status_code == 403

    def test_financial_can_access(self, http_fin):
        r = http_fin.get(PREFIX + "/simulate", params={"reference_month": REF_MONTH})
        assert r.status_code == 200

    def test_unauthenticated_returns_401(self, http_unauth):
        r = http_unauth.get(PREFIX + "/simulate", params={"reference_month": REF_MONTH})
        assert r.status_code == 401


# ---------------------------------------------------------------------------
# GET /simulate/pdf
# ---------------------------------------------------------------------------

class TestSimulatePdf:
    def test_returns_200(self, http):
        r = http.get(PREFIX + "/simulate/pdf", params={"reference_month": REF_MONTH})
        assert r.status_code == 200

    def test_content_type_is_pdf(self, http):
        r = http.get(PREFIX + "/simulate/pdf", params={"reference_month": REF_MONTH})
        assert "application/pdf" in r.headers.get("content-type", "")

    def test_pdf_has_content(self, http):
        r = http.get(PREFIX + "/simulate/pdf", params={"reference_month": REF_MONTH})
        assert len(r.content) > 0

    def test_pdf_starts_with_pdf_header(self, http):
        r = http.get(PREFIX + "/simulate/pdf", params={"reference_month": REF_MONTH})
        assert r.content[:4] == b"%PDF"

    def test_requires_reference_month(self, http):
        r = http.get(PREFIX + "/simulate/pdf")
        assert r.status_code == 422

    def test_filter_client_without_client_id_returns_422(self, http):
        r = http.get(PREFIX + "/simulate/pdf", params={
            "reference_month": REF_MONTH,
            "filter_type": "client",
        })
        assert r.status_code == 422

    def test_operational_cannot_access(self, http_op):
        r = http_op.get(PREFIX + "/simulate/pdf", params={"reference_month": REF_MONTH})
        assert r.status_code == 403

    def test_financial_can_access(self, http_fin):
        r = http_fin.get(PREFIX + "/simulate/pdf", params={"reference_month": REF_MONTH})
        assert r.status_code == 200

    def test_unauthenticated_returns_401(self, http_unauth):
        r = http_unauth.get(PREFIX + "/simulate/pdf", params={"reference_month": REF_MONTH})
        assert r.status_code == 401


class TestSimulateXlsx:
    def test_filter_client_without_client_id_returns_422(self, http):
        r = http.get(PREFIX + "/simulate/xlsx", params={
            "reference_month": REF_MONTH,
            "filter_type": "client",
        })
        assert r.status_code == 422


# ---------------------------------------------------------------------------
# POST /generate
# ---------------------------------------------------------------------------

class TestGenerate:
    def test_forwards_exact_selections_for_every_category(self, http, monkeypatch):
        import app.api.v1.endpoints.billing_closure as endpoint

        captured = {}

        def fake_execute(db, ref, filter_type, client_id, **kwargs):
            captured.update(kwargs)
            return {
                'reference_month': '05/2025', 'generated': 0,
                'billing_ids': [], 'consolidated_unico': 0, 'total_amount': 0,
                'uninstall_fees_generated': 0, 'uninstall_events_processed': 0,
                'uninstall_fees_deferred': 0, 'uninstall_fees_skipped': 0,
                'uninstall_billing_ids': [], 'services_generated': 0,
                'service_billing_ids': [], 'total_services_amount': 0,
                'grand_total': 0,
            }

        monkeypatch.setattr(endpoint, 'execute_closure', fake_execute)
        r = http.post(PREFIX + '/generate', params=[
            ('reference_month', REF_MONTH),
            ('contract_ids', '10'), ('contract_ids', '11'),
            ('uninstall_event_ids', '20'), ('uninstall_event_ids', '21'),
            ('charge_item_ids', '30'), ('charge_item_ids', '31'),
        ])

        assert r.status_code == 200
        assert captured == {
            'contract_ids': [10, 11],
            'uninstall_event_ids': [20, 21],
            'charge_item_ids': [30, 31],
        }

    def test_returns_200(self, http):
        r = http.post(PREFIX + "/generate", params={"reference_month": REF_MONTH})
        assert r.status_code == 200

    def test_status_is_completed(self, http):
        r = http.post(PREFIX + "/generate", params={"reference_month": REF_MONTH})
        assert r.json()["status"] == "completed"

    def test_response_echoes_reference_month(self, http):
        r = http.post(PREFIX + "/generate", params={"reference_month": REF_MONTH})
        assert r.json()["reference_month"] == REF_MONTH

    def test_reference_month_da_resposta_e_aceito_de_volta(self, http):
        """Round-trip: o mês devolvido tem que ser reenviável à própria API.

        O `**result` do serviço sobrescrevia o eco do parâmetro e a resposta
        saía como 05/2025 — formato que o endpoint rejeita com 422.
        """
        primeira = http.post(PREFIX + "/generate", params={"reference_month": REF_MONTH})
        devolvido = primeira.json()["reference_month"]

        segunda = http.post(PREFIX + "/generate", params={"reference_month": devolvido})
        assert segunda.status_code == 200

    def test_response_traz_mes_formatado_para_exibicao(self, http):
        r = http.post(PREFIX + "/generate", params={"reference_month": REF_MONTH})
        assert r.json()["reference_month_label"] == "05/2025"

    def test_response_has_result_fields(self, http):
        r = http.post(PREFIX + "/generate", params={"reference_month": REF_MONTH})
        data = r.json()
        for key in (
            "status", "reference_month", "generated", "total_amount",
            "uninstall_fees_generated", "uninstall_fees_skipped",
            "services_generated", "total_services_amount", "grand_total",
        ):
            assert key in data, f"missing key: {key}"

    def test_generated_is_integer(self, http):
        r = http.post(PREFIX + "/generate", params={"reference_month": REF_MONTH})
        assert isinstance(r.json()["generated"], int)

    def test_grand_total_is_numeric(self, http):
        r = http.post(PREFIX + "/generate", params={"reference_month": REF_MONTH})
        assert isinstance(r.json()["grand_total"], (int, float))

    def test_requires_reference_month(self, http):
        r = http.post(PREFIX + "/generate")
        assert r.status_code == 422

    def test_filter_client_without_client_id_returns_422(self, http):
        r = http.post(PREFIX + "/generate", params={
            "reference_month": REF_MONTH,
            "filter_type": "client",
        })
        assert r.status_code == 422

    def test_second_call_generates_nothing(self, http, db):
        first = http.post(PREFIX + "/generate", params={"reference_month": REF_MONTH})
        r2 = http.post(PREFIX + "/generate", params={"reference_month": REF_MONTH})
        assert r2.status_code == 200
        assert r2.json()["generated"] == 0
        first_id = first.json()['closure_batch_id']
        second_id = r2.json()['closure_batch_id']
        assert first_id != second_id
        assert db.get(ClosureJob, second_id).result['payment_billing_ids'] == []
        assert {first_id, second_id} <= {row['id'] for row in http.get(PREFIX + '/lotes').json()}
        preview = http.get(f'{PREFIX}/lotes/{second_id}')
        assert preview.status_code == 200
        assert preview.json()['itens'] == []

    def test_operational_cannot_generate(self, http_op):
        r = http_op.post(PREFIX + "/generate", params={"reference_month": REF_MONTH})
        assert r.status_code == 403

    def test_financial_can_generate(self, http_fin):
        r = http_fin.post(PREFIX + "/generate", params={"reference_month": REF_MONTH})
        assert r.status_code == 200
        assert r.json()["status"] == "completed"

    def test_unauthenticated_returns_401(self, http_unauth):
        r = http_unauth.post(PREFIX + "/generate", params={"reference_month": REF_MONTH})
        assert r.status_code == 401
